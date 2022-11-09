from __future__ import print_function

import numpy as np
import pandas as pd
import re
from openpyxl import load_workbook


class MultiLang:
    COLUMN_KEY = 'Key'
    COLUMN_LANG_SUPPORTED = 'languages supported'
    PATH_ANDROID_BASE = "../app/src/main/res/values-"
    PATH_LOCAL_FILE = '../translation/Translation String.xlsx'  # path to file + file name
    SHEET_NAME = 'Multi'

    LANGUAGE_SUPPORTED = []
    COLUMNS_TO_PARSE = []

    def __init__(self):
        return

    def read_file_content(self):
        file_path = self.PATH_LOCAL_FILE

        wb = load_workbook(file_path)
        content = pd.DataFrame()

        for sheet in wb.worksheets:
            # write sheet's names that you want to use,
            if not sheet.title.startswith(self.SHEET_NAME):
                continue

            values = sheet.values
            data_frame = pd.DataFrame(values)
            content = content.append(data_frame, ignore_index=True)
            if not values:
                print('No data found.')
            # else:
            #     print (rangeName[i])
            #     print (content)

        content = self.clean_content(content)

        # parse the language supported from Sheet
        for i in range(len(content)):
            lang = str(content[self.COLUMN_LANG_SUPPORTED][i])
            if lang != "" and lang != 'None' and lang is not None:
                self.LANGUAGE_SUPPORTED.append(lang)

        # load the columns that will be parsed
        for langCode in self.LANGUAGE_SUPPORTED:
            # accessibility_column = "%s Accessibility" % langCode
            self.COLUMNS_TO_PARSE.append(langCode)
            # columns_to_parse.append(accessibility_column)

        # print('---------------------------------')
        # print('=================================')
        content = self.clear_brackets(content)

        return content.copy(deep=True)

    '''This function cleans the content dataframe and names columns for convenient use'''

    def clean_content(self, data_frame):
        # Name columns by first row of copy sheet
        data_frame = data_frame.rename(columns=data_frame.iloc[0]).drop(data_frame.index[0])
        # Remove columns without a name
        data_frame = data_frame[
            [col for col in data_frame.columns if col is not None and col is not np.nan]]
        # Drop duplicates keys and rows with "None", only whitespace, or "" value in COLUMN_KEY column
        data_frame = data_frame.drop_duplicates(subset=self.COLUMN_KEY, keep='first')
        data_frame = data_frame[data_frame[self.COLUMN_KEY].notnull()]

        data_frame = data_frame.drop(
            data_frame[data_frame[self.COLUMN_KEY].str.strip() == ''].index)
        # Drop rows with "key" value in COLUMN_KEY column and reset index (first row of each worksheet)
        data_frame = data_frame.drop(
            data_frame[data_frame[self.COLUMN_KEY] == self.COLUMN_KEY].index).reset_index(drop=True)

        # since some keys are written as `key.key`, make sure to replace `.` with `_`
        for item in range(len(data_frame[self.COLUMN_KEY])):
            if "." in data_frame[self.COLUMN_KEY][item]:
                data_frame[self.COLUMN_KEY][item] = data_frame[self.COLUMN_KEY][item].replace(".","_")
                data_frame[self.COLUMN_KEY][item] = data_frame[self.COLUMN_KEY][item].replace("&","_")

        return data_frame

    '''Replace all [variable placeholders] to []'''

    def clear_brackets(self, df):
        for col in range(len(self.COLUMNS_TO_PARSE)):
            for item in range(len(df[self.COLUMNS_TO_PARSE[col]])):
                df[self.COLUMNS_TO_PARSE[col]][item] = re.sub("([\[]).*?([\]])", "[]", self.if_null(
                    df[self.COLUMNS_TO_PARSE[col]][item]))
                df[self.COLUMNS_TO_PARSE[col]][item] = re.sub(r"\n", " ", self.if_null(
                    df[self.COLUMNS_TO_PARSE[col]][item]))
        return df

    def if_null(self, item):
        if item:
            return str(item)
        else:
            x = ''
            return str(x)
# from MultiLang import MultiLang

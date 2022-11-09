from __future__ import print_function
from openpyxl import load_workbook

import os
import numpy as np
import pandas as pd
import re
import time
from googleapiclient.discovery import build
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from httplib2 import Http
from oauth2client import file, client, tools

#this is needed for Google script to login 1st time, TODO clean later

# If modifying these scopes, delete the file token.json. Token file is generated automatically after a successful login for the script
SCOPES = ['https://www.googleapis.com/auth/spreadsheets.readonly']

# Credentials path relative to fastlane execution
PATH_TOKEN = '../translation/token.json'
PATH_CREDENTIALS = '../translation/credentials.json'

# The ID and range of a sample spreadsheet.
# SAMPLE_RANGE_NAME = 'Translation'
# our script
# https://docs.google.com/spreadsheets/d/{docId}/edit#gid=1556281741
SPREADSHEET_ID = '18P9iZY0zqnxsqZejW-JXGOqysrjbFwn5aXf6YXip5zY'

LANGUAGE_SUPPORTED = []
COLUMNS_TO_PARSE = []


class GoogleSheet:
    COLUMN_KEY = 'Key'
    COLUMN_LANG_SUPPORTED = 'languages supported'
    SHEET_NAME = 'Multi'

    LANGUAGE_SUPPORTED = []
    COLUMNS_TO_PARSE = []

    def __init__(self):
        return

    def getContent(self):
        # The file token.json stores the user's access and refresh tokens, and is
        # created automatically when the authorization flow completes for the first time.
        store = file.Storage(PATH_TOKEN)
        creds = store.get()

        if not creds or creds.invalid:
            flow = client.flow_from_clientsecrets(PATH_CREDENTIALS, scope=SCOPES)
            creds = tools.run_flow(flow, store)

        service = build('sheets', 'v4', http=creds.authorize(Http()))

        # Call the# else:
        #         #     print (rangeName[i])
        #         #     print (content) Sheets API
        sheets = service.spreadsheets()
        # Create DataFrame to load content into
        content = pd.DataFrame()
        # Iterate through all worksheets identified in the range and load them into the content dataframe
        for sheet in sheets.get(spreadsheetId=SPREADSHEET_ID).execute().get('sheets', ''):
            sheetTitle = sheet.get('properties', {}).get('title')

            # write sheet's names that you want to use,
            if (not sheetTitle.startswith(self.SHEET_NAME)) \
                    and (not sheetTitle.startswith('Date Formats')):
                continue

            result = sheets.values().get(spreadsheetId=SPREADSHEET_ID, range=sheetTitle).execute()
            values = result.get('values', [])
            df = pd.DataFrame(values)
            content = content.append(df, ignore_index=True)
            if not values:
                print('No data found.')
            # else:
            #     print (rangeName[i])
            #     print (content)

        content = self.clean_content(content)

        # parse the language supported from Sheet
        for i in range(len(content)):
            lang = str(content[self.COLUMN_LANG_SUPPORTED][i])
            if lang != "" and lang != 'None' and lang is not None:
                self.LANGUAGE_SUPPORTED.append(lang)

        # load the columns that will be parsed
        for langCode in self.LANGUAGE_SUPPORTED:
            # accessibility_column = "%s Accessibility" % langCode
            self.COLUMNS_TO_PARSE.append(langCode)
            # columns_to_parse.append(accessibility_column)

        # print('---------------------------------')
        # print('=================================')
        content = self.clear_brackets(content)

        return content.copy(deep=True)

    '''This function cleans the content dataframe and names columns for convenient use'''

    def clean_content(self, data_frame):
        # Name columns by first row of copy sheet
        data_frame = data_frame.rename(columns=data_frame.iloc[0]).drop(data_frame.index[0])
        # Remove columns without a name
        data_frame = data_frame[
            [col for col in data_frame.columns if col is not None and col is not np.nan]]
        # Drop duplicates keys and rows with "None", only whitespace, or "" value in COLUMN_KEY column
        data_frame = data_frame.drop_duplicates(subset=self.COLUMN_KEY, keep='first')
        data_frame = data_frame[data_frame[self.COLUMN_KEY].notnull()]

        data_frame = data_frame.drop(
            data_frame[data_frame[self.COLUMN_KEY].str.strip() == ''].index)
        # Drop rows with "key" value in COLUMN_KEY column and reset index (first row of each worksheet)
        data_frame = data_frame.drop(
            data_frame[data_frame[self.COLUMN_KEY] == self.COLUMN_KEY].index).reset_index(drop=True)

        # since some keys are written as `key.key`, make sure to replace `.` with `_`
        for item in range(len(data_frame[self.COLUMN_KEY])):
            if "." in data_frame[self.COLUMN_KEY][item]:
                data_frame[self.COLUMN_KEY][item] = data_frame[self.COLUMN_KEY][item].replace(".","_")
                data_frame[self.COLUMN_KEY][item] = data_frame[self.COLUMN_KEY][item].replace("&","_")

        return data_frame

    '''Replace all [variable placeholders] to []'''

    '''Replace all [variable placeholders] to []'''

    def clear_brackets(self, df):
        for col in range(len(self.COLUMNS_TO_PARSE)):
            for item in range(len(df[self.COLUMNS_TO_PARSE[col]])):
                df[self.COLUMNS_TO_PARSE[col]][item] = re.sub("([\[]).*?([\]])", "[]", self.if_null(
                    df[self.COLUMNS_TO_PARSE[col]][item]))
                df[self.COLUMNS_TO_PARSE[col]][item] = re.sub(r"\n", " ", self.if_null(
                    df[self.COLUMNS_TO_PARSE[col]][item]))
        return df

    def if_null(self, item):
        if item:
            return str(item)
        else:
            x = ''
            return str(x)

# def main():
#     """Shows basic usage of the Sheets API.
#     Prints values from a sample spreadsheet.
#     """
#     creds = None
#     # The file token.json stores the user's access and refresh tokens, and is
#     # created automatically when the authorization flow completes for the first
#     # time.
#     if os.path.exists('token.json'):
#         creds = Credentials.from_authorized_user_file('token.json', SCOPES)
#     # If there are no (valid) credentials available, let the user log in.
#     if not creds or not creds.valid:
#         if creds and creds.expired and creds.refresh_token:
#             creds.refresh(Request())
#         else:
#             flow = InstalledAppFlow.from_client_secrets_file(
#                 'credentials.json', SCOPES)
#             creds = flow.run_local_server(port=0)
#         # Save the credentials for the next run
#         with open('token.json', 'w') as token:
#             token.write(creds.to_json())
#
#     service = build('sheets', 'v4', credentials=creds)
#
#     # Call the Sheets API
#     sheet = service.spreadsheets()
#     result = sheet.values().get(spreadsheetId=SAMPLE_SPREADSHEET_ID,
#                                 range=SAMPLE_RANGE_NAME).execute()
#     values = result.get('values', [])
#
#     if not values:
#         print('No data found.')
#     else:
#         print('Name, Major:')
#         for row in values:
#             # Print columns A and E, which correspond to indices 0 and 4.
#             print('%s, %s' % (row[0], row[1]))
#         content = pd.DataFrame()
#         df=pd.DataFrame(values)
#         content = content.append(df, ignore_index=True)
#
# if __name__ == '__main__':
#     main()
